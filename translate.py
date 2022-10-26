import requests
from docx.api import Document

def translate(word):
    url = 'https://fanyi.baidu.com/sug'
    headers = {'User-Agent':'Mozilla/5.0 (Macintosh; Intel Mac OS X 10_9_5)'
    'AppleWebKit 537.36 (KHTML, like Gecko) Chrome',
    'Accept':'text/html,application/xhtml+xml,application/xml;'
    'q=0.9,image/webp,*/*;q=0.8'}
    Form_data = {'kw': word}
    response = requests.post(url, data=Form_data,headers=headers)
    if len(response.json()['data'])!=0:
        return str(response.json()['data'][0]['v'])
    else:
        return "Not found."

document = Document('test.docx')
table = document.tables[0]
i= 0
while i < len(table.rows):
        text = str(table.cell(i,0).text)
        print(translate(text))
        table.cell(i,1).text = translate(text)
        i = i + 1
document.save('test.docx')
